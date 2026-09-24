#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
cubicaje.py — compara los m3 reclamados por viaje contra la lista de cubicaje
por placa (TERRAPLEN.xlsx, hoja "Cubicaje 2026") o, para GRANULARES, contra
el simple registro de placas existentes (GRANULARES.xlsx, hoja "Placas"),
ya que en GRANULARES el volumen sale de báscula y no se compara con lista
salvo que la placa no exista.

Reglas (diseño en diseno_conciliador_completo.md, punto 4):
  - ambito == "TERRAPLEN": compara m3 del viaje contra Volumen m3 de
    "Cubicaje 2026" para esa placa. Tolerancia 0.05 (valor absoluto de la
    diferencia). Si la placa no aparece en la lista -> discrepancia
    "placa sin cubicaje".
  - ambito == "GRANULARES": NO se compara m3 (viene de báscula). Solo se
    avisa si la placa no existe en la hoja "Placas" de GRANULARES.xlsx.
  - ambito distinto/desconocido (None, "AMBAS", etc.): se reporta como
    "ambito_desconocido" (no se puede decidir la regla) y no se compara.

USO:
  python cubicaje.py --terraplen TERRAPLEN.xlsx --granulares GRANULARES.xlsx \
      --viajes viajes.json [--json]

  viajes.json: lista de objetos {"remision":..., "placa":..., "m3":..., "ambito":...}
"""

import argparse
import json
import sys

try:
    import openpyxl
except ImportError:
    print("Falta la dependencia 'openpyxl' (pip install openpyxl).", file=sys.stderr)
    sys.exit(2)

TOLERANCIA_M3 = 0.05

HOJA_CUBICAJE_TERRAPLEN = "Cubicaje 2026"
HOJA_PLACAS_GRANULARES = "Placas"


def _norm_placa(p):
    if p is None:
        return ""
    return str(p).strip().upper()


def cargar_cubicaje_terraplen(ruta_terraplen):
    """Placa (normalizada) -> volumen m3 (float), desde 'Cubicaje 2026'.

    Columnas esperadas (fila 1 encabezado): N°, PlacaVolqueta, Volumen m3
    (Formato nuevo), Fecha de Cubicaje, Fecha de Re Cubicaje, Observación.
    Se buscan por nombre de encabezado (tolerando acentos/espacios), no por
    posición fija, salvo fallback a B/C si no calzan los nombres.
    """
    wb = openpyxl.load_workbook(ruta_terraplen, data_only=True, read_only=True)
    if HOJA_CUBICAJE_TERRAPLEN not in wb.sheetnames:
        raise ValueError(
            "TERRAPLEN.xlsx no tiene la hoja '%s' (hojas: %s)"
            % (HOJA_CUBICAJE_TERRAPLEN, ", ".join(wb.sheetnames))
        )
    ws = wb[HOJA_CUBICAJE_TERRAPLEN]
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return {}
    encabezados = [str(h or "").strip().lower() for h in filas[0]]

    def _col(*claves):
        for i, h in enumerate(encabezados):
            for clave in claves:
                if clave in h:
                    return i
        return None

    col_placa = _col("placa")
    col_vol = _col("volumen")
    if col_placa is None:
        col_placa = 1  # columna B
    if col_vol is None:
        col_vol = 2  # columna C

    resultado = {}
    for fila in filas[1:]:
        if col_placa >= len(fila):
            continue
        placa = _norm_placa(fila[col_placa])
        if not placa:
            continue
        vol = fila[col_vol] if col_vol < len(fila) else None
        try:
            vol = float(vol)
        except (TypeError, ValueError):
            continue
        resultado[placa] = vol
    return resultado


def cargar_placas_granulares(ruta_granulares):
    """Conjunto de placas registradas en GRANULARES.xlsx, hoja 'Placas'
    (columnas PLACA, EMPRESA)."""
    wb = openpyxl.load_workbook(ruta_granulares, data_only=True, read_only=True)
    if HOJA_PLACAS_GRANULARES not in wb.sheetnames:
        raise ValueError(
            "GRANULARES.xlsx no tiene la hoja '%s' (hojas: %s)"
            % (HOJA_PLACAS_GRANULARES, ", ".join(wb.sheetnames))
        )
    ws = wb[HOJA_PLACAS_GRANULARES]
    filas = list(ws.iter_rows(values_only=True))
    placas = set()
    for fila in filas[1:]:
        if not fila:
            continue
        placa = _norm_placa(fila[0])
        if placa:
            placas.add(placa)
    return placas


def _num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def comparar(viajes, cubicaje_terraplen, placas_granulares):
    """Devuelve la lista de discrepancias (dicts) para los viajes dados."""
    discrepancias = []
    for viaje in viajes:
        placa = _norm_placa(viaje.get("placa"))
        ambito = (viaje.get("ambito") or "").strip().upper()
        remision = viaje.get("remision")
        m3 = _num(viaje.get("m3"))

        if ambito == "TERRAPLEN":
            if placa not in cubicaje_terraplen:
                discrepancias.append({
                    "remision": remision,
                    "placa": placa,
                    "ambito": ambito,
                    "tipo": "placa_sin_cubicaje",
                    "detalle": "Placa no está en TERRAPLEN 'Cubicaje 2026'",
                })
                continue
            vol = cubicaje_terraplen[placa]
            if m3 is None:
                discrepancias.append({
                    "remision": remision,
                    "placa": placa,
                    "ambito": ambito,
                    "tipo": "m3_invalido",
                    "detalle": "El viaje no trae m3 numérico",
                })
                continue
            diferencia = round(m3 - vol, 4)
            if abs(diferencia) > TOLERANCIA_M3:
                discrepancias.append({
                    "remision": remision,
                    "placa": placa,
                    "ambito": ambito,
                    "tipo": "m3_distinto_cubicaje",
                    "m3_viaje": m3,
                    "m3_cubicaje": vol,
                    "diferencia": diferencia,
                })
        elif ambito == "GRANULARES":
            # Báscula: no se compara m3, solo existencia de placa.
            if placa not in placas_granulares:
                discrepancias.append({
                    "remision": remision,
                    "placa": placa,
                    "ambito": ambito,
                    "tipo": "placa_sin_cubicaje",
                    "detalle": "Placa no está en GRANULARES 'Placas'",
                })
        else:
            discrepancias.append({
                "remision": remision,
                "placa": placa,
                "ambito": viaje.get("ambito"),
                "tipo": "ambito_desconocido",
                "detalle": "No se puede decidir la regla de cubicaje sin ámbito TERRAPLEN/GRANULARES",
            })
    return discrepancias


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--terraplen", required=True, help="Ruta a TERRAPLEN.xlsx")
    ap.add_argument("--granulares", required=True, help="Ruta a GRANULARES.xlsx")
    ap.add_argument("--viajes", required=True, help="JSON con la lista de viajes")
    ap.add_argument("--json", action="store_true", help="Salida en JSON")
    args = ap.parse_args()

    with open(args.viajes, "r", encoding="utf-8") as f:
        viajes = json.load(f)

    cubicaje_terraplen = cargar_cubicaje_terraplen(args.terraplen)
    placas_granulares = cargar_placas_granulares(args.granulares)

    discrepancias = comparar(viajes, cubicaje_terraplen, placas_granulares)

    resumen = {
        "total_viajes": len(viajes),
        "total_discrepancias": len(discrepancias),
        "placas_cubicaje_terraplen": len(cubicaje_terraplen),
        "placas_registradas_granulares": len(placas_granulares),
        "discrepancias": discrepancias,
    }

    if args.json:
        print(json.dumps(resumen, indent=2, ensure_ascii=False))
    else:
        print(
            "Viajes: %d | Discrepancias: %d | Cubicaje TERRAPLEN: %d placas | Placas GRANULARES: %d"
            % (
                resumen["total_viajes"],
                resumen["total_discrepancias"],
                resumen["placas_cubicaje_terraplen"],
                resumen["placas_registradas_granulares"],
            )
        )
        for d in discrepancias:
            print(" - [%s] rem=%s placa=%s %s: %s" % (
                d["ambito"], d.get("remision"), d.get("placa"), d["tipo"],
                d.get("detalle", d),
            ))

    sys.exit(1 if discrepancias else 0)


if __name__ == "__main__":
    main()
