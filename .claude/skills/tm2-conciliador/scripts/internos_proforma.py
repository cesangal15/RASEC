"""tm2-conciliador · internos_proforma.py

Detecta los días internos que la proforma cobra SIN decir «interno» (sep-2026, Suministros 2.ª Q sep:
TFT906 21/09, 7 viajes de 5,7 km con la cuota del día en una celda combinada de V.total):

  python internos_proforma.py --proforma "<PROFORMA>.xlsx" [--out internos_proforma.json]

Por cada hoja busca:
- celdas COMBINADAS en la columna de valor total (V.total / Total) que abarcan varias filas (la cuota
  de la placa ese día en una sola celda);
- las filas de una hoja que se llame «INTERNOS» (la cobra el contratista aparte).
La tarifa (V.unitario) NO es señal: depende de los km del viaje (0–3, 3–5, > 5 km), no del día interno.
y agrupa por placa + fecha. Es solo una señal: la decisión (día interno, CC, observación) la confirma
el dueño (`references/internos_y_cuota.md`). Solo lee la proforma; no escribe nada salvo --out.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
import unicodedata
from collections import defaultdict

import openpyxl


def norm(v) -> str:
    s = unicodedata.normalize("NFKD", str(v or "")).encode("ascii", "ignore").decode().upper()
    return " ".join(s.replace(".", " ").split())


def columnas(ws):
    """Fila de encabezados y columnas (1-based) de fecha, remisión, placa, V.unitario y V.total.
    Los formatos cambian por contratista («No. Remisión» / «No. RECIBO» / «RECIBO»; «V.unitario» /
    «PRECIO» / «VLR M3/KM» / «VLR DIA DE TRABAJO»; «V.total» / «TOTAL» / «VLR TOTAL TRANSPORTE»)."""
    for r in range(1, min(ws.max_row, 30) + 1):
        hdr = {c: norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)}
        rem = next((c for c, h in hdr.items() if "REMIS" in h), None) or             next((c for c, h in hdr.items() if "RECIBO" in h and "CLIENTE" not in h), None)
        if not rem:
            continue
        fecha = next((c for c, h in hdr.items() if h.startswith("FECHA")), None)
        placa = next((c for c, h in hdr.items() if "PLACA" in h), None)
        vunit = next((c for c, h in hdr.items() if "UNITARIO" in h or h == "PRECIO"
                      or h.startswith("VLR M3") or h.startswith("VLR DIA")), None)
        vtot = next((c for c, h in hdr.items() if h in ("V TOTAL", "TOTAL", "VALOR TOTAL", "VLR TOTAL TRANSPORTE")), None)
        return r, fecha, rem, placa, vunit, vtot
    return None


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--proforma", required=True)
    ap.add_argument("--out")
    a = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")

    wb = openpyxl.load_workbook(a.proforma)            # combinadas
    wv = openpyxl.load_workbook(a.proforma, data_only=True)  # valores
    hallazgos = []
    for ws in wb.worksheets:
        col = columnas(ws)
        if not col:
            print(f"(hoja «{ws.title.strip()}»: sin columna de remisión reconocible, no se revisó)")
            continue
        h, cf, cr, cp, cu, ct = col
        wsv = wv[ws.title]
        # solo filas de viaje: remisión y fecha de verdad (debajo a veces hay dinámicas o totales)
        filas = [r for r in range(h + 1, ws.max_row + 1) if wsv.cell(r, cr).value not in (None, "")
                 and (not cf or isinstance(wsv.cell(r, cf).value, (dt.datetime, dt.date)))]
        hoja_interna = "INTERN" in norm(ws.title)
        marcas = defaultdict(lambda: {"filas": [], "motivos": set(), "valor": None})

        def clave(r):
            f = wsv.cell(r, cf).value if cf else None
            f = f.date().isoformat() if isinstance(f, dt.datetime) else str(f)
            return (f, str(wsv.cell(r, cp).value if cp else ""))

        for m in ws.merged_cells.ranges:
            if ct and m.min_row > h and m.min_col <= ct <= m.max_col and m.max_row > m.min_row:
                for r in range(m.min_row, m.max_row + 1):
                    if r in filas:
                        g = marcas[clave(r)]
                        g["filas"].append(r); g["motivos"].add(f"V.total combinado {m.coord}")
                        g["valor"] = wsv.cell(m.min_row, m.min_col).value
        for r in filas:
            if hoja_interna:
                g = marcas[clave(r)]
                g["filas"].append(r); g["motivos"].add("hoja de INTERNOS")
        for (fecha, placa), g in sorted(marcas.items()):
            rs = sorted(set(g["filas"]))
            hallazgos.append({"hoja": ws.title, "fecha": fecha, "placa": placa,
                              "remisiones": [str(wsv.cell(r, cr).value) for r in rs],
                              "valor_combinado": g["valor"], "motivos": sorted(g["motivos"])})

    for x in hallazgos:
        print(f"{x['hoja'].strip()} · {x['fecha']} · {x['placa']} · {len(x['remisiones'])} viaje(s) "
              f"{', '.join(x['remisiones'])} · {'; '.join(x['motivos'])}"
              + (f" · valor {x['valor_combinado']}" if x["valor_combinado"] is not None else ""))
    if not hallazgos:
        print("Sin señales de día interno cobrado sin marcar (ni V.total combinado ni hoja INTERNOS).")
    if a.out:
        json.dump(hallazgos, open(a.out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return 0


if __name__ == "__main__":
    sys.exit(main())
